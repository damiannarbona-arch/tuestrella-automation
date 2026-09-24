from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active; ws.title = "Calculadora"
F = "Arial"
BLUE = Font(name=F, color="0000FF"); BLACK = Font(name=F); BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14); HDR = Font(name=F, bold=True, color="FFFFFF")
YEL = PatternFill("solid", fgColor="FFFF00"); HDRF = PatternFill("solid", fgColor="3F5147")
SEC = PatternFill("solid", fgColor="E9E4DA"); RES = PatternFill("solid", fgColor="DDEBDD")
thin = Side(style="thin", color="BBBBBB"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
EUR = '#,##0.00 €;(#,##0.00 €);-'; PCT = '0.0%;(0.0%);-'; NUM = '#,##0.0;(#,##0.0);-'

ws["A1"] = "Kivoa · Calculadora de margen por pedido (importes en €)"; ws["A1"].font = TITLE
ws["A2"] = "Celdas en AZUL con fondo AMARILLO = datos a rellenar/ajustar. Celdas en negro = fórmulas (no tocar). Los valores de ejemplo son SUPUESTOS: sustitúyelos por tus precios reales de proveedor, plan de Shopify y publicidad."
ws["A2"].font = Font(name=F, italic=True, size=9); ws.merge_cells("A2:F2"); ws["A2"].alignment = Alignment(wrap_text=True); ws.row_dimensions[2].height = 30

cols = ["Concepto", "A · Gelato (UE) → cliente España/UE", "B · Gelato (imprenta EE. UU.) → cliente EE. UU.", "C · Proveedor China → cliente EE. UU.", "D · Proveedor China → cliente UE", "Nota"]
r = 4
for i, c in enumerate(cols, 1):
    cell = ws.cell(row=r, column=i, value=c); cell.font = HDR; cell.fill = HDRF; cell.alignment = Alignment(wrap_text=True, vertical="center"); cell.border = BOX
ws.row_dimensions[r].height = 45

rows = {}
def section(title):
    global r
    r += 1
    ws.cell(row=r, column=1, value=title).font = BOLD
    for c in range(1, 7): ws.cell(row=r, column=c).fill = SEC

def inp(key, label, vals, fmt, note=""):
    global r
    r += 1; rows[key] = r
    ws.cell(row=r, column=1, value=label).font = BLACK
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=2 + j, value=v); c.font = BLUE; c.fill = YEL; c.number_format = fmt; c.border = BOX
    ws.cell(row=r, column=6, value=note).font = Font(name=F, size=9, italic=True)
    ws.cell(row=r, column=6).alignment = Alignment(wrap_text=True)

def calc(key, label, tmpl, fmt, bold=False, fill=None, note=""):
    global r
    r += 1; rows[key] = r
    ws.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
    for j in range(4):
        col = get_column_letter(2 + j)
        f = tmpl.format(c=col, **{k: v for k, v in rows.items()})
        c = ws.cell(row=r, column=2 + j, value=f); c.font = BOLD if bold else BLACK; c.number_format = fmt; c.border = BOX
        if fill: c.fill = fill
    ws.cell(row=r, column=6, value=note).font = Font(name=F, size=9, italic=True)
    ws.cell(row=r, column=6).alignment = Alignment(wrap_text=True)

section("1 · Venta")
inp("precio", "Precio que paga el cliente (€)", [59.90, 74.90, 74.90, 59.90], EUR,
    "UE: precio con IVA incluido. EE. UU.: precio sin sales tax (allí el impuesto se suma aparte en el checkout). Ejemplo EE. UU.: 79,99 $ ≈ 74,90 €.")
inp("iva", "IVA incluido en el precio (%)", [0.21, 0, 0, 0.21], PCT,
    "España 21 %. Ventas a otros países UE: 21 % hasta 10.000 €/año de ventas UE transfronterizas; después, IVA del país del cliente vía OSS. Venta a EE. UU.: sin IVA español.")

section("2 · Coste del producto (proveedor)")
inp("coste", "Coste del producto sin IVA (€)", [26.00, 28.00, 9.00, 9.00], EUR,
    "SUPUESTO. Pon el precio real de Gelato (catálogo, 30×40 enmarcado) o la cotización del proveedor chino. El IVA que te cobre un proveedor español/UE es deducible si estás en régimen general.")
inp("envio", "Envío proveedor → cliente (€)", [7.00, 8.50, 9.00, 8.00], EUR, "SUPUESTO. China: envío con seguimiento de 10 a 20 días.")
inp("arancel_pct", "Arancel (% sobre valor del producto)", [0, 0, 0.30, 0], PCT,
    "Gelato imprime en el país del cliente, así que no hay aduana. China → EE. UU.: aproximadamente 25–35 % (Sección 301 + otros; el de minimis está suspendido). Verifica la partida exacta.")
inp("arancel_fijo", "Arancel fijo por artículo (€)", [0, 0, 0, 3.00], EUR,
    "China → UE: 3 € por artículo desde el 1/7/2026 (sustituye la exención de 150 €). Temporal hasta julio de 2028.")
inp("aduana", "Gestión aduanera / despacho por paquete (€)", [0, 0, 5.00, 2.00], EUR,
    "SUPUESTO. Tarifa del transitario o de la paquetería por el despacho formal.")
inp("iva_import", "IVA de importación no recuperable (€)", [0, 0, 0, 0], EUR,
    "China → UE con IOSS: el IVA de importación se cobra al cliente dentro del precio (fila IVA). Déjalo en 0 salvo que el cliente pague IVA en destino y lo absorbas tú.")

section("3 · Costes de venta")
inp("pasarela_pct", "Comisión pasarela de pago (%)", [0.019, 0.034, 0.034, 0.019], PCT,
    "SUPUESTO. Revisa tu plan en Shopify > Configuración > Pagos. Las tarjetas extranjeras y la conversión de divisa suben el %.")
inp("pasarela_fijo", "Comisión fija por transacción (€)", [0.25, 0.30, 0.30, 0.25], EUR, "SUPUESTO.")
inp("cac", "Publicidad por pedido / CAC (€)", [15.00, 20.00, 20.00, 15.00], EUR,
    "Gasto en anuncios dividido entre pedidos. Suele ser el coste más grande: mídelo cada mes.")
inp("incid", "Incidencias: reimpresiones y reembolsos (% del coste producto + envío)", [0.03, 0.03, 0.10, 0.10], PCT,
    "SUPUESTO. Con China, los envíos largos y la calidad variable suben las roturas y reclamaciones.")
inp("otros", "Otros por pedido: packaging, apps prorrateadas (€)", [1.00, 1.00, 1.00, 1.00], EUR, "")

section("4 · Resultado por pedido")
calc("neto", "Ingreso neto sin IVA (€)", "={c}{precio}/(1+{c}{iva})", EUR)
calc("iva_eur", "IVA a ingresar a Hacienda (€)", "={c}{precio}-{c}{neto}", EUR)
calc("aranceles", "Aranceles + aduana (€)", "={c}{coste}*{c}arancel_pct_+{c}{arancel_fijo}+{c}{aduana}+{c}{iva_import}".replace("{c}arancel_pct_", "{c}{arancel_pct}"), EUR)
calc("pasarela", "Comisión pasarela (€)", "={c}{precio}*{c}{pasarela_pct}+{c}{pasarela_fijo}", EUR)
calc("incid_eur", "Coste incidencias (€)", "=({c}{coste}+{c}{envio})*{c}{incid}", EUR)
calc("costes", "Total costes variables (€)", "={c}{coste}+{c}{envio}+{c}{aranceles}+{c}{pasarela}+{c}{incid_eur}+{c}{cac}+{c}{otros}", EUR, bold=True)
calc("benef", "BENEFICIO POR PEDIDO antes de impuesto sobre la renta (€)", "={c}{neto}-{c}{costes}", EUR, bold=True, fill=RES)
calc("margen", "Margen neto sobre ingreso sin IVA (%)", "=IF({c}{neto}=0,0,{c}{benef}/{c}{neto})", PCT, bold=True, fill=RES)
calc("margen_sin_ads", "Margen de contribución SIN publicidad (%)", "=IF({c}{neto}=0,0,({c}{benef}+{c}{cac})/{c}{neto})", PCT,
     note="Lo que queda antes de pagar anuncios: te dice cuánto puedes invertir en captar cada cliente.")

section("5 · Precio mínimo recomendado")
ws.cell(row=r+1, column=1, value="Margen neto objetivo (%)").font = BLACK
r += 1; rows["objetivo"] = r
c = ws.cell(row=r, column=2, value=0.20); c.font = BLUE; c.fill = YEL; c.number_format = PCT; c.border = BOX
ws.cell(row=r, column=6, value="Se aplica a los 4 escenarios. Por debajo del 15–20 % cualquier imprevisto (anuncios más caros, una reimpresión) se come el beneficio.").font = Font(name=F, size=9, italic=True)
ws.cell(row=r, column=6).alignment = Alignment(wrap_text=True)
calc("costes_fijos_pedido", "Costes por pedido que no dependen del precio (€)", "={c}{costes}-{c}{precio}*{c}{pasarela_pct}", EUR)
calc("pmin", "PRECIO MÍNIMO de venta para el margen objetivo (€, en formato del escenario)",
     "=IF(((1-$B${objetivo})/(1+{c}{iva})-{c}{pasarela_pct})<=0,\"n/a\",{c}{costes_fijos_pedido}/((1-$B${objetivo})/(1+{c}{iva})-{c}{pasarela_pct}))",
     EUR, bold=True, fill=RES, note="Fórmula: costes fijos por pedido ÷ [(1 − margen) ÷ (1 + IVA) − % pasarela]. Compara con tu precio actual (fila 1).")

section("6 · Visión mensual (se aplica a cada escenario por separado)")
inp("pedidos", "Pedidos al mes", [20, 20, 20, 20], NUM, "SUPUESTO. Cámbialo para ver el punto de equilibrio.")
inp("fijos", "Costes fijos mensuales (€): Shopify + apps + cuota autónomo + gestoría", [180, 180, 180, 180], EUR,
    "SUPUESTO: Shopify Basic ~36 € + apps ~20 € + cuota de autónomo (tarifa plana 80 € el primer año; después según ingresos) + gestoría ~45 €.")
inp("irpf", "Tipo efectivo IRPF estimado sobre el beneficio (%)", [0.20, 0.20, 0.20, 0.20], PCT,
    "Como residente fiscal en España tributas en IRPF por TODO lo que ganes, venda donde venda (también con una LLC de EE. UU.). Tu tipo depende del resto de ingresos (nómina): pide el dato a tu gestor.")
calc("benef_mes", "Beneficio mensual antes de IRPF (€)", "={c}{benef}*{c}{pedidos}-{c}{fijos}", EUR, bold=True)
calc("irpf_eur", "IRPF estimado (€)", "=MAX(0,{c}{benef_mes})*{c}{irpf}", EUR)
calc("neto_mes", "BENEFICIO NETO MENSUAL en tu bolsillo (€)", "={c}{benef_mes}-{c}{irpf_eur}", EUR, bold=True, fill=RES)
calc("equilibrio", "Pedidos/mes para cubrir costes fijos (punto de equilibrio)", "=IF({c}{benef}<=0,\"no rentable\",{c}{fijos}/{c}{benef})", NUM, bold=True)

ws.column_dimensions["A"].width = 58
for col in "BCDE": ws.column_dimensions[col].width = 22
ws.column_dimensions["F"].width = 70
ws.freeze_panes = "B5"
for row in ws.iter_rows(min_row=5, max_row=r):
    for c in row:
        if c.column == 1: c.alignment = Alignment(wrap_text=True, vertical="center")

# --- Comparativa
c2 = wb.create_sheet("Comparativa")
data = [
 ["Factor", "A · Gelato → España/UE", "B · Gelato (EE. UU.) → EE. UU.", "C · China → EE. UU.", "D · China → UE"],
 ["Dónde se fabrica", "Imprenta de la red Gelato en la UE", "Imprenta Gelato en EE. UU.", "Fábrica en China", "Fábrica en China"],
 ["Plazo total al cliente", "6–12 días", "5–10 días", "12–25 días + aduana", "10–20 días + aduana"],
 ["Aranceles", "No", "No (producción local)", "Sí: ~25–35 % + despacho formal por paquete", "Sí: 3 € por artículo desde 7/2026"],
 ["IVA / impuesto de venta", "IVA 21 % (OSS si superas 10.000 €/año de ventas a otros países UE)", "Sin IVA español. Sales tax de EE. UU. solo si superas los umbrales de cada estado (~100.000 $/año); en Etsy lo gestiona Etsy", "Igual que B", "IVA del país del cliente vía IOSS"],
 ["Impuesto sobre la renta", "IRPF en España", "IRPF en España", "IRPF en España", "IRPF en España"],
 ["Personalización con foto", "Sí, con aprobación manual", "Sí, con aprobación manual", "Posible pero lenta; errores caros de corregir", "Igual que C"],
 ["Calidad / control", "Alta y constante", "Alta", "Variable: pide muestras", "Variable"],
 ["Riesgo de reclamaciones", "Bajo", "Bajo", "Alto (plazos, roturas, aduana)", "Medio-alto"],
 ["Capital inmovilizado", "0 € (bajo pedido)", "0 €", "0 € si es dropshipping; stock si compras al por mayor", "Igual que C"],
 ["Mejor para", "Validar el negocio ya", "Escalar con margen en dólares", "Productos genéricos sin personalización", "Productos genéricos baratos"],
]
for i, row in enumerate(data, 1):
    for j, v in enumerate(row, 1):
        cell = c2.cell(row=i, column=j, value=v); cell.font = HDR if i == 1 else Font(name=F, bold=(j == 1))
        if i == 1: cell.fill = HDRF
        cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.border = BOX
c2.column_dimensions["A"].width = 26
for col in "BCDE": c2.column_dimensions[col].width = 34

# --- Fuentes
s = wb.create_sheet("Fuentes")
src = [
 ["Dato", "Fuente"],
 ["EE. UU.: fin del de minimis de 800 $ para todos los países (29/08/2025)", "https://www.easyship.com/blog/de-minimis-exemption-suspension"],
 ["EE. UU.: aranceles a China tras la sentencia del Supremo (feb. 2026), Sección 301", "https://www.china-briefing.com/news/supreme-court-tariff-ruling-china-impact/"],
 ["EE. UU.–UE: techo arancelario del 15 % desde el 1/7/2026", "https://www.tariffstool.com/guides/eu-us-trade-deal-15-percent-tariff-live"],
 ["UE: fin de la exención de 150 €; 3 € por artículo desde el 1/7/2026", "https://www.consilium.europa.eu/en/press/press-releases/2025/12/12/customs-council-agrees-to-levy-customs-duty-on-small-parcels-as-of-1-july-2026/"],
 ["Gelato: producción local en el país del cliente", "https://www.gelato.com/print-on-demand"],
 ["IVA en tienda online, OSS, autónomo", "https://gestoria247.com/blog/impuestos-tienda-online-2026"],
 ["Nota", "Los importes de coste, envío, comisiones y publicidad son SUPUESTOS de ejemplo; sustitúyelos por tus datos reales. Esto no es asesoramiento fiscal: valida IRPF, OSS y sales tax con tu gestor."],
]
for i, row in enumerate(src, 1):
    for j, v in enumerate(row, 1):
        cell = s.cell(row=i, column=j, value=v); cell.font = HDR if i == 1 else Font(name=F)
        if i == 1: cell.fill = HDRF
        cell.alignment = Alignment(wrap_text=True, vertical="top")
s.column_dimensions["A"].width = 60; s.column_dimensions["B"].width = 90
wb.save("calculadora-margenes-kivoa.xlsx")
print(rows)

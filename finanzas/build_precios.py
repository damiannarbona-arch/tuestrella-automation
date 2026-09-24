from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
F="Arial"
BLUE=Font(name=F,color="0000FF"); BLACK=Font(name=F); BOLD=Font(name=F,bold=True)
HDR=Font(name=F,bold=True,color="FFFFFF"); HDRF=PatternFill("solid",fgColor="3F5147")
YEL=PatternFill("solid",fgColor="FFFF00"); RES=PatternFill("solid",fgColor="DDEBDD"); SEC=PatternFill("solid",fgColor="E9E4DA")
thin=Side(style="thin",color="BBBBBB"); BOX=Border(top=thin,bottom=thin,left=thin,right=thin)
EUR='#,##0.00 €;[Red](#,##0.00 €);-'; USD='[$$-409]#,##0.00;[Red]([$$-409]#,##0.00);-'; PCT='0.0%;[Red](0.0%);-'

wb=Workbook(); ws=wb.active; ws.title="Precios por país"
ws["A1"]="Kivoa · Tabla de precios por país y canal (retrato enmarcado, Gelato estándar)"; ws["A1"].font=Font(name=F,bold=True,size=14)
ws["A2"]="Amarillo/azul = datos que cambias tú. Negro = fórmulas. Costes de Gelato del 24/09/2026 para España y EE. UU. (captura del usuario, sin Gelato+, se toma el marco más caro de cada tamaño). Los demás países: rellena desde Gelato → 'Previsualiza tus costes para…'."
ws["A2"].font=Font(name=F,italic=True,size=9); ws.merge_cells("A2:T2"); ws["A2"].alignment=Alignment(wrap_text=True); ws.row_dimensions[2].height=30

# Global parameters
ws["A4"]="Parámetros globales"; ws["A4"].font=BOLD
params=[("Margen objetivo antes de IRPF (sobre ingreso sin IVA)",0.20,PCT,"Por debajo del 15–20 % cualquier imprevisto se come el beneficio."),
        ("Tipo efectivo IRPF estimado",0.20,PCT,"Depende de tus otros ingresos. Pídelo a tu gestor."),
        ("Tipo de cambio: 1 USD = … EUR",0.86,'0.0000',"SUPUESTO. Actualízalo el día que fijes precios."),
        ("Comisión fija pasarela por pedido (€)",0.25,EUR,"Shopify Payments: revisa tu plan.")]
P={}
for i,(lab,val,fmt,note) in enumerate(params):
    r=5+i; ws.cell(row=r,column=1,value=lab).font=BLACK
    c=ws.cell(row=r,column=2,value=val); c.font=BLUE; c.fill=YEL; c.number_format=fmt; c.border=BOX
    ws.cell(row=r,column=3,value=note).font=Font(name=F,size=9,italic=True)
    P[i]=f"$B${r}"
M,IRPF,FX,FIJO=P[0],P[1],P[2],P[3]

hdr=["Canal / país","Tamaño","Coste producto Gelato (€ sin IVA)","Envío Gelato (€)","Coste total (€)",
     "PRECIO VENTA ACTUAL (€, con IVA si aplica)","IVA del país (%)","Ingreso sin IVA (€)","IVA a pagar (€)",
     "Comisión % (pasarela / marketplace)","Comisiones (€)","Publicidad por pedido (€)",
     "Beneficio antes IRPF (€)","IRPF (€)","BENEFICIO NETO (€)","Margen neto (%)",
     "Precio mínimo para margen objetivo (€)","PRECIO RECOMENDADO (€)","Precio recomendado (US$)","Nota"]
HR=10
for j,h in enumerate(hdr,1):
    c=ws.cell(row=HR,column=j,value=h); c.font=HDR; c.fill=HDRF; c.alignment=Alignment(wrap_text=True,vertical="center"); c.border=BOX
ws.row_dimensions[HR].height=60

# name, size, cost, ship, price, vat, fee%, ads, note, usd?
ES=[("20×25 cm",13.32,6.63,39.90),("30×40 cm",20.84,6.63,59.90),("50×70 cm",41.39,19.30,99.90)]
US=[("8×10\"",17.70,11.72,58.84),("12×16\"",28.33,11.72,80.10),("20×28\"",49.80,25.21,150.02)]
blocks=[
 ("Web · España",ES,0.21,0.019,12,False,"Costes reales Gelato → España. IVA 21 %."),
 ("Web · EE. UU.",US,0.0,0.034,15,True,"Costes reales Gelato → EE. UU. Sin IVA; el sales tax se suma aparte. Comisión mayor por tarjeta extranjera + cambio de divisa."),
 ("Etsy · España",ES,0.21,0.1122,0,False,"Etsy: 6,5 % transacción + 4 % pago + 0,72 % tasa regulatoria (España). Además 0,20 $ por anuncio (incluido en el fijo). Sin publicidad propia: Etsy trae el tráfico. Si activas Offsite Ads, suma 15 %."),
 ("Etsy · EE. UU.",US,0.0,0.1122,0,True,"Igual que Etsy España. En EE. UU., Etsy cobra y paga el sales tax por ti."),
 ("Web · Alemania",[("20×25 cm",None,None,59.90),("30×40 cm",None,None,74.90),("50×70 cm",None,None,119.90)],0.19,0.019,12,False,"RELLENA costes Gelato → Alemania. IVA 19 % (aplica cuando superes 10.000 €/año de ventas a otros países UE; hasta entonces, 21 % español)."),
 ("Web · Francia",[("20×25 cm",None,None,59.90),("30×40 cm",None,None,74.90),("50×70 cm",None,None,119.90)],0.20,0.019,12,False,"RELLENA costes Gelato → Francia. IVA 20 % (misma regla de los 10.000 €)."),
 ("Web · Italia",[("20×25 cm",None,None,59.90),("30×40 cm",None,None,74.90),("50×70 cm",None,None,119.90)],0.22,0.019,12,False,"RELLENA costes Gelato → Italia. IVA 22 % (misma regla de los 10.000 €)."),
]
r=HR
for name,items,vat,fee,ads,usd,note in blocks:
    r+=1
    for c in range(1,21): ws.cell(row=r,column=c).fill=SEC
    ws.cell(row=r,column=1,value=name).font=BOLD
    ws.cell(row=r,column=20,value=note).font=Font(name=F,size=9,italic=True); ws.cell(row=r,column=20).alignment=Alignment(wrap_text=True)
    ws.row_dimensions[r].height=42
    etsy = name.startswith("Etsy")
    for size,cost,ship,price in items:
        r+=1
        ws.cell(row=r,column=1,value=name).font=BLACK; ws.cell(row=r,column=2,value=size).font=BLACK
        for col,val,fmt in [(3,cost,EUR),(4,ship,EUR),(6,price,EUR),(7,vat,PCT),(10,fee,PCT),(12,ads,EUR)]:
            c=ws.cell(row=r,column=col,value=val); c.font=BLUE; c.fill=YEL; c.number_format=fmt; c.border=BOX
        fixed = f"({FIJO}+0.18)" if etsy else FIJO
        f={5:f"=IF(OR(C{r}=\"\",D{r}=\"\"),\"\",C{r}+D{r})",
           8:f"=F{r}/(1+G{r})",9:f"=F{r}-H{r}",
           11:f"=F{r}*J{r}+{fixed}",
           13:f"=IF(E{r}=\"\",\"\",H{r}-E{r}-K{r}-L{r})",
           14:f"=IF(M{r}=\"\",\"\",MAX(0,M{r})*{IRPF})",
           15:f"=IF(M{r}=\"\",\"\",M{r}-N{r})",
           16:f"=IF(OR(M{r}=\"\",H{r}=0),\"\",O{r}/H{r})",
           17:f"=IF(E{r}=\"\",\"\",(E{r}+{fixed}+L{r})/((1-{M})/(1+G{r})-J{r}))",
           18:f"=IF(Q{r}=\"\",\"\",CEILING(Q{r},5)-0.1)",
           19:(f"=IF(Q{r}=\"\",\"\",CEILING(Q{r}/{FX},10)-0.01)" if usd else "")}
        for col,formula in f.items():
            c=ws.cell(row=r,column=col,value=formula if formula else None); c.border=BOX
            c.number_format = USD if col==19 else (PCT if col==16 else EUR)
            c.font = BOLD if col in (15,18,19) else BLACK
            if col in (15,18,19): c.fill=RES
widths=[16,11,14,11,11,15,9,12,10,12,11,12,13,10,13,10,14,14,14,60]
for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=w
ws.freeze_panes="C11"

# Explanation sheet
e=wb.create_sheet("Cómo leerla")
txt=[
 "CÓMO SE CALCULA (por pedido)",
 "1. Ingreso sin IVA = precio ÷ (1 + IVA). En EE. UU. el IVA es 0: el precio es ingreso íntegro.",
 "2. Beneficio antes de IRPF = ingreso sin IVA − (producto + envío Gelato) − comisiones − publicidad por pedido.",
 "3. IRPF = beneficio × tu tipo efectivo. Beneficio neto = lo que te queda de verdad.",
 "4. Precio mínimo = (coste total + comisión fija + publicidad) ÷ [(1 − margen objetivo) ÷ (1 + IVA) − % comisión].",
 "5. Precio recomendado = precio mínimo redondeado hacia arriba a un precio comercial (…4,90 / …9,90 € y …9,99 $).",
 "",
 "POR QUÉ NO FIARTE DEL 'BENEFICIO ESTIMADO' DE GELATO",
 "• Gelato NO descuenta el IVA: en España, de un cuadro de 39,90 €, 6,92 € son de Hacienda, no tuyos.",
 "• En tu captura de España tenías DESMARCADO 'Incluir el costo de envío': por eso te salía un 66 % de margen. Con IVA y envío, el 20×25 a 39,90 € deja 13 € antes de comisiones y publicidad.",
 "• Tampoco incluye comisiones de pago, publicidad, reimpresiones ni tu IRPF.",
 "",
 "LECTURA RÁPIDA DE TUS PRECIOS ACTUALES (con 12 € de publicidad por pedido)",
 "• 20×25 a 39,90 €: beneficio prácticamente 0. Súbelo o úsalo solo como producto de entrada sin anuncios (Etsy).",
 "• 50×70 a 99,90 €: el envío de 19,30 € se come el margen. Necesita rondar 115 €.",
 "• Etsy sin publicidad es el canal más rentable al principio, pese a su comisión del 11 %.",
 "",
 "AVISOS",
 "• Costes de Gelato sin IVA: si eres autónomo en régimen general, el IVA que te cobra Gelato es deducible. Pregunta a tu gestor si en tu caso aplicaría el recargo de equivalencia (comercio minorista), porque cambia este cálculo.",
 "• Otros países de la UE: hasta 10.000 €/año de ventas a otros países UE se aplica el IVA español (21 %). Por encima, el del país del cliente vía OSS.",
 "• Reino Unido: vender allí puede obligarte a registrarte en el IVA británico desde la primera venta. No lo actives sin gestor.",
 "• Publicidad por pedido: es el dato que más cambia. Actualízalo cada mes con lo que realmente gastas ÷ pedidos.",
 "• Esto no es asesoramiento fiscal: valida IRPF, IVA y OSS con tu gestor.",
]
for i,t in enumerate(txt,1):
    c=e.cell(row=i,column=1,value=t); c.font=BOLD if t.isupper() or t.startswith(("CÓMO","POR QUÉ","LECTURA","AVISOS")) else Font(name=F)
    c.alignment=Alignment(wrap_text=True)
e.column_dimensions["A"].width=130
wb.save("tabla-precios-kivoa.xlsx"); print("ok", r)
